"""Slide 3: Industry Pipeline Breakdown - horizontal bar chart."""

from pptx.presentation import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

from src.deck.styles import (
    DARK_CHARCOAL, WHITE_RABBIT, PURPLE_RAIN, BLUE_MOON,
    YELLOW_SUBMARINE, MID_GRAY, CARD_BG_DARK,
    FONT_HEADING, FONT_BODY,
)
from src.deck.slides.helpers import (
    add_blank_slide, set_slide_bg, add_textbox,
)
from src.data.models import ReportData
from src.insights.metrics import format_currency

# Bar colors — brand palette first, then complementary extras for overflow.
# All chosen to be visible against DarkCharcoal (#15161e).
BAR_COLORS = [
    PURPLE_RAIN,                        # #6559FF
    BLUE_MOON,                          # #00D1FF
    YELLOW_SUBMARINE,                   # #FFD700
    RGBColor(0x00, 0xC8, 0x53),         # Success Green
    RGBColor(0xFF, 0xA0, 0x00),         # Warm Amber
    RGBColor(0xE0, 0x6C, 0xFF),         # Soft Violet
    RGBColor(0x00, 0xE6, 0xA0),         # Mint
    RGBColor(0xFF, 0x6B, 0x6B),         # Coral
    RGBColor(0x7C, 0xB3, 0xFF),         # Sky Blue
    RGBColor(0xC8, 0xE6, 0x00),         # Lime
]


def build(prs: Presentation, data: ReportData, metrics: dict, config=None, **kwargs):
    from src.config import ReportConfig
    if config is None:
        config = ReportConfig()

    slide = add_blank_slide(prs)
    set_slide_bg(slide, DARK_CHARCOAL)

    # Title
    add_textbox(
        slide, "TA Pipeline by Industry",
        Inches(0.75), Inches(0.4), Inches(6), Inches(0.5),
        font_name=FONT_HEADING, font_size=Pt(24),
        font_color=WHITE_RABBIT, bold=True,
    )

    # Subtitle with total pipeline figure
    add_textbox(
        slide,
        f"{metrics['open_pipeline_aud_fmt']} across {metrics['open_pipeline_count']} active deals",
        Inches(0.75), Inches(0.95), Inches(8), Inches(0.3),
        font_name=FONT_BODY, font_size=Pt(12),
        font_color=MID_GRAY,
    )

    # Sort industries by total AUD descending
    sorted_industries = sorted(
        data.industry_open_pipeline, key=lambda x: x.total_aud, reverse=True
    )

    # Filter out zero-value industries
    sorted_industries = [ind for ind in sorted_industries if ind.total_aud > 0]

    if not sorted_industries:
        add_textbox(
            slide, "No open pipeline data available",
            Inches(0.75), Inches(3.5), Inches(6), Inches(0.5),
            font_name=FONT_BODY, font_size=Pt(14),
            font_color=MID_GRAY,
        )
        return

    max_aud = sorted_industries[0].total_aud if sorted_industries else 1

    # Chart layout — fits within 7.5" slide height
    # Available vertical space: 1.5" (start) to ~7.2" (bottom margin) = 5.7"
    label_left = Inches(0.75)
    label_width = Inches(2.6)
    chart_left = Inches(3.5)        # Where bars start (after labels)
    max_bar_right = Inches(10.5)    # Max bar extent (leaves room for value labels)
    max_bar_width = max_bar_right - chart_left
    start_y = Inches(1.5)

    # Cap at 10 industries, dynamically size rows to fit
    display_industries = sorted_industries[:10]
    num_rows = len(display_industries)
    available_height = Inches(5.7)
    row_height = int(available_height / num_rows)
    bar_height = int(row_height * 0.65)
    bar_pad = int((row_height - bar_height) / 2)

    for i, ind in enumerate(display_industries):
        y = start_y + i * row_height
        color = BAR_COLORS[i % len(BAR_COLORS)]

        # Industry label (right-aligned, vertically centered in row)
        add_textbox(
            slide, ind.industry_type,
            label_left, y, label_width, row_height,
            font_name=FONT_BODY, font_size=Pt(11),
            font_color=WHITE_RABBIT, bold=False,
            alignment=PP_ALIGN.RIGHT,
        )

        # Bar width proportional to max value
        bar_ratio = ind.total_aud / max_aud if max_aud > 0 else 0
        bar_width = int(max_bar_width * bar_ratio)
        # Minimum visible bar width
        bar_width = max(bar_width, Inches(0.3))

        # Draw the bar
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            chart_left, y + bar_pad,
            bar_width, bar_height,
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
        bar.adjustments[0] = 0.15

        # Value label (to the right of the bar, vertically centered)
        value_text = f"{format_currency(ind.total_aud)}  ({ind.total_count} deals)"
        add_textbox(
            slide, value_text,
            chart_left + bar_width + Inches(0.15), y,
            Inches(2.5), row_height,
            font_name=FONT_BODY, font_size=Pt(10),
            font_color=MID_GRAY,
        )
